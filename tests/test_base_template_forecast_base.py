"""The BASE template's historical block, and the DCF that depends on it.

Two defects, one cause. `build_template.py` styled the IS historical block
BLACK -- this repository's convention for "same-sheet formula" -- although
its own comment calls those cells "hardcoded actuals the user enters". So
42 real input cells per instance were invisible to every tool that
identifies an input by its blue font, and could stay empty with nothing
reporting a gap.

`IS!F5 = E5*(1+Assumptions!C5)`, so an empty `IS!E5` zeroes every projected
year, the cash flow statement, and the whole DCF. That does not surface as
a zero valuation. `Equity value = Enterprise value - net debt`, and net debt
is often negative, so the committed Microsoft FY2024 corporate-finance case
published an implied value per share of **$1.13** from an enterprise value
of exactly 0. Zero is a tell a reader notices; $1.13 is not.

The tests below are weighted toward the discrimination that matters: the
guard must fire on a dead model AND must get out of the way the moment a
forecast base exists. A guard that always says "n/a" would pass a naive
version of the first test and be worthless.
"""
from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

from recalc import recalc

from tools.verify_template_exhaustion import INPUT_FONT_RGB, _rgb

ROOT = Path(__file__).resolve().parents[1]
BASE_TEMPLATES = (
    ROOT / "01_Investment_Banking" / "_template_BASE.xlsx",
    ROOT / "02_Corporate_Finance" / "_template_BASE.xlsx",
    ROOT / "12_Equity_Finance" / "_template_BASE.xlsx",
)
DEAD_CASE = ROOT / "02_Corporate_Finance" / "instances" / "public_microsoft_2024.xlsx"
HISTORICAL_ROWS = range(5, 21)
HISTORICAL_COLUMNS = range(3, 6)  # C, D, E -- FY-2A, FY-1A, FY0A


# _rgb and INPUT_FONT_RGB are imported from the scanner rather than
# restated here. openpyxl's aRGB alpha byte is not stable across how a
# Font was constructed -- "000000FF" and "FF0000FF" are the same blue --
# and the scanner already normalises that (it is why domain 31 once
# reported zero candidate cells). A second copy of the rule in the test
# would agree with the scanner today and drift from it later, and a test
# that disagrees with the tool it is testing is worse than no test.
def _cell_rgb(cell) -> str | None:
    return _rgb(cell.font.color if cell.font else None)


def _historical_inputs(sheet) -> list[str]:
    """Cells in the historical block that carry no formula, i.e. inputs."""
    found = []
    for row in HISTORICAL_ROWS:
        for column in HISTORICAL_COLUMNS:
            cell = sheet.cell(row=row, column=column)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                continue
            found.append(cell.coordinate)
    return found


class HistoricalBlockVisibilityTests(unittest.TestCase):
    def test_every_historical_input_carries_the_input_font(self) -> None:
        for template in BASE_TEMPLATES:
            with self.subTest(template=template.name, domain=template.parent.name):
                sheet = openpyxl.load_workbook(template)["IS"]
                inputs = _historical_inputs(sheet)
                # Control: if the block ever stops being input-shaped this
                # test must fail loudly rather than pass over an empty set.
                self.assertGreaterEqual(
                    len(inputs), 40, "the historical block has stopped offering inputs"
                )
                unmarked = [c for c in inputs if _cell_rgb(sheet[c]) != INPUT_FONT_RGB]
                self.assertEqual(
                    [], unmarked,
                    f"{len(unmarked)} historical input cells are invisible to every "
                    "tool that identifies an input by its blue font",
                )

    def test_derived_cells_in_the_block_are_not_marked_as_inputs(self) -> None:
        """The marker must mean something. If the rule were 'paint the whole
        block blue', a formula cell would be called an input and the coverage
        denominator would be wrong in the other direction."""
        for template in BASE_TEMPLATES:
            with self.subTest(template=template.name):
                sheet = openpyxl.load_workbook(template)["IS"]
                formulas = [
                    sheet.cell(row=r, column=c).coordinate
                    for r in HISTORICAL_ROWS
                    for c in HISTORICAL_COLUMNS
                    if isinstance(sheet.cell(row=r, column=c).value, str)
                    and str(sheet.cell(row=r, column=c).value).startswith("=")
                ]
                self.assertGreater(len(formulas), 0, "no derived cells to discriminate against")
                mismarked = [c for c in formulas if _cell_rgb(sheet[c]) == INPUT_FONT_RGB]
                self.assertEqual([], mismarked)

    def test_the_coverage_scanner_now_sees_them(self) -> None:
        """End to end through the repository's own scanner, not a reimplementation."""
        from tools import verify_template_exhaustion as scanner

        candidates = scanner.find_candidate_cells(BASE_TEMPLATES[1])
        historical = {
            coord for coord in candidates.get("IS", set())
            if coord[0] in "CDE" and coord[1:].isdigit() and 5 <= int(coord[1:]) <= 20
        }
        self.assertGreaterEqual(len(historical), 40)


class ForecastBaseGuardTests(unittest.TestCase):
    """The published per-share output, recalculated for real."""

    @classmethod
    def setUpClass(cls) -> None:
        cls._tmp = tempfile.TemporaryDirectory()
        tmp = Path(cls._tmp.name)

        cls.dead = tmp / "dead.xlsx"
        shutil.copyfile(DEAD_CASE, cls.dead)
        result = recalc(str(cls.dead), timeout=90)
        assert result.get("status") == "success" and not result.get("total_errors", 1), result
        cls.dead_values = openpyxl.load_workbook(cls.dead, data_only=True)

        # The same workbook with a forecast base supplied. A deliberately
        # synthetic figure: this test is about the mechanism, not about
        # Microsoft, and a real revenue here would look like a sourced fact.
        cls.live = tmp / "live.xlsx"
        shutil.copyfile(DEAD_CASE, cls.live)
        book = openpyxl.load_workbook(cls.live)
        book["IS"]["E5"] = 1000.0
        book["IS"]["E7"] = 300.0
        book.save(cls.live)
        result = recalc(str(cls.live), timeout=90)
        assert result.get("status") == "success" and not result.get("total_errors", 1), result
        cls.live_values = openpyxl.load_workbook(cls.live, data_only=True)

    @classmethod
    def tearDownClass(cls) -> None:
        cls._tmp.cleanup()

    def test_the_fixture_really_has_no_forecast_base(self) -> None:
        """A control. Without it the guard tests could pass because the
        fixture changed rather than because the guard works."""
        base = openpyxl.load_workbook(DEAD_CASE)["IS"]["E5"].value
        self.assertIn(base, (None, "", 0), f"fixture now has a forecast base: {base!r}")

    def test_a_dcf_with_no_base_publishes_no_price(self) -> None:
        dcf = self.dead_values["DCF"]
        self.assertNotIsInstance(
            dcf["I14"].value, (int, float),
            "a per-share price was published from a model with no forecast base",
        )
        self.assertIn("no forecast base", str(dcf["I14"].value))

    def test_the_absence_is_stated_rather_than_left_to_be_inferred(self) -> None:
        self.assertIn("MISSING", str(self.dead_values["DCF"]["I15"].value))

    def test_the_arithmetic_that_produced_the_published_number_is_still_there(self) -> None:
        """Pin the mechanism, not just the symptom.

        Enterprise value 0 with negative net debt still yields a positive
        equity value -- that is a real property of the bridge and this
        change does not pretend otherwise. What changed is that it is no
        longer divided by shares and published as a price.
        """
        dcf = self.dead_values["DCF"]
        self.assertEqual(0, dcf["I10"].value)
        self.assertLess(dcf["I11"].value, 0)
        self.assertGreater(dcf["I12"].value, 0)

    def test_supplying_a_base_restores_a_real_price(self) -> None:
        """The discrimination that makes the guard a guard.

        A blanket 'n/a' would satisfy every test above and be useless.
        """
        dcf = self.live_values["DCF"]
        self.assertIn("present", str(dcf["I15"].value))
        self.assertIsInstance(dcf["I14"].value, (int, float))
        self.assertNotEqual(0, dcf["I10"].value)

    def test_a_live_case_built_before_this_change_is_untouched(self) -> None:
        """The one BASE-template workbook with a real forecast base keeps
        its valuation, so this is a guard on absent models rather than a
        change to how a working DCF values anything."""
        proxy = ROOT / "01_Investment_Banking" / "instances" / "public_adobe_dcf_proxy.xlsx"
        if not proxy.exists():
            self.skipTest("agent-tool draft not present")
        dcf = openpyxl.load_workbook(proxy, data_only=True)["DCF"]
        self.assertIsInstance(dcf["I14"].value, (int, float))
        self.assertGreater(dcf["I14"].value, 0)


if __name__ == "__main__":
    unittest.main()


class AgentToolGuardTests(unittest.TestCase):
    """tools/agents/dcf_comps.py rejected a non-numeric or negative price
    and a non-numeric enterprise value. It did not reject an enterprise
    value of exactly zero, which is how $1.13 got past it."""

    def test_a_zero_enterprise_value_is_rejected(self) -> None:
        from tools.agents import dcf_comps

        errors = dcf_comps._check_workbook_structurally_sound(DEAD_CASE)
        self.assertTrue(
            any("exactly 0" in error for error in errors),
            f"a model with no forecast base passed the structural check: {errors}",
        )

    def test_a_real_valuation_still_passes(self) -> None:
        """Discrimination. A check that rejects everything is not a check."""
        from tools.agents import dcf_comps

        proxy = ROOT / "01_Investment_Banking" / "instances" / "public_adobe_dcf_proxy.xlsx"
        if not proxy.exists():
            self.skipTest("agent-tool draft not present")
        self.assertEqual([], dcf_comps._check_workbook_structurally_sound(proxy))


class DcfMemoRefusalTests(unittest.TestCase):
    """The deck builder renders DCF!I14 with a currency format.

    Now that the cell is text when there is no forecast base, rendering a
    dead model would fail as "Unknown format code 'f' for object of type
    'str'" -- a correct refusal wearing a name that sends the reader to
    python's formatter instead of to IS!E5. It refuses on its own terms
    instead, and the message names the cell to fill.

    Driven through the real build_dcf_memo with a synthetic manifest rather
    than by writing a fixture manifest into the evidence directory.
    """

    def _manifest_for(self, output_rel: str) -> dict:
        return {
            "output": output_rel,
            "as_of": "2026-01-01",
            "cover": {"Title:": "Fixture Co -- test"},
            "inputs": [],
        }

    def test_a_dead_model_refuses_with_the_cell_to_fill(self) -> None:
        from unittest.mock import patch

        from tools.builders import build_dcf_memo as memo

        dead = str(DEAD_CASE.relative_to(ROOT))
        with patch.object(memo, "_load_manifest", return_value=self._manifest_for(dead)):
            with self.assertRaises(SystemExit) as caught:
                memo.build_dcf_memo("fixture", Path(tempfile.mkdtemp()) / "out.pptx")
        message = str(caught.exception)
        self.assertIn("no valuation to render", message)
        self.assertIn("IS!E5", message)

    def test_a_live_model_still_renders(self) -> None:
        """Discrimination again: the guard must only stop dead models."""
        from tools.builders import build_dcf_memo as memo

        proxy = ROOT / "01_Investment_Banking" / "instances" / "public_adobe_dcf_proxy.xlsx"
        if not (proxy.parent / "public_adobe_dcf_proxy.manifest.json").exists():
            self.skipTest("agent-tool draft manifest not present")
        out = Path(tempfile.mkdtemp()) / "live.pptx"
        memo.build_dcf_memo("public_adobe_dcf_proxy", out)
        self.assertTrue(out.exists())
