"""Tests for the base-year reproduction gate.

The gate's whole claim is that it can tell a model that agrees with a known
period from one that does not, so the tests that matter are the ones that make
it say both things. Two of them exist because the tool got them wrong first:

  * `test_a_workbook_whose_only_actual_is_the_anchor_is_not_a_pass` -- the first
    version compared the revenue row, which this probe HOLDS equal to the
    disclosed figure. One real case (the Adobe DCF proxy) sources revenue and
    nothing else, and it reported `reproduced` on a comparison of the probe with
    itself. A check that cannot fail is not a check.

  * `test_a_text_result_fails_rather_than_passes` -- an Excel line that cannot
    compute lands on `"-"`, and text sorts above every number. Anything that
    compares without asking whether it got a number first will read that as fine.
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tools"))

from tools import verify_base_year_reproduction as gate  # noqa: E402


def build_workbook(path: Path, *, actuals: dict[int, float] | None = None,
                   opex_driver: float = 0.25, include_is: bool = True) -> Path:
    """A minimal BASE-shaped statement model whose chain is sound by construction.

    Derived, not copied from a real case: a fixture pointed at whichever
    artefact happens to be broken today stops being a fixture when it is fixed.
    """
    book = openpyxl.Workbook()
    assumptions = book.active
    assumptions.title = "Assumptions"
    assumptions["C5"] = 0.10           # revenue growth
    assumptions["C6"] = 0.40           # COGS share of revenue
    assumptions["C7"] = opex_driver    # opex share of revenue
    if not include_is:
        book.save(path)
        return path
    sheet = book.create_sheet("IS")
    sheet["B4"], sheet["C4"], sheet["D4"] = "Line item", "FY-1A", "FY0A"
    sheet["E4"] = "FY1E"
    sheet["B5"], sheet["B7"], sheet["B8"], sheet["B10"] = "Revenue", "COGS", "Gross profit", "Opex"
    sheet["B11"] = "EBIT"
    for row, value in (actuals or {}).items():
        sheet.cell(row=row, column=4).value = value
    sheet["E5"] = "=D5*(1+Assumptions!C5)"
    sheet["E7"] = "=E5*Assumptions!C6"
    sheet["E8"] = "=E5-E7"
    sheet["E10"] = "=E5*Assumptions!C7"
    sheet["E11"] = "=E8-E10"
    book.save(path)
    return path


SOUND_ACTUALS = {5: 1000.0, 7: 400.0, 8: 600.0, 10: 250.0, 11: 350.0}


class ComparisonTests(unittest.TestCase):
    """The verdict logic, with no spreadsheet and no LibreOffice involved."""

    def _compare(self, disclosed, reproduced, exemptions=None, tolerance=0.5):
        labels = {row: f"line {row}" for row in disclosed}
        return {
            line["row"]: line
            for line in gate.compare(disclosed, reproduced, labels,
                                     exemptions or {}, tolerance)
        }

    def test_a_line_within_tolerance_reproduces(self):
        lines = self._compare({14: 100.0}, {14: 100.4})
        self.assertEqual(lines[14]["verdict"], "reproduced")

    def test_a_line_outside_tolerance_deviates(self):
        lines = self._compare({14: 109433.0}, {14: 84970.71})
        self.assertEqual(lines[14]["verdict"], "deviates")
        self.assertAlmostEqual(lines[14]["delta"], -24462.29, places=2)
        self.assertAlmostEqual(lines[14]["delta_pct"], -22.35, places=2)

    def test_a_text_result_fails_rather_than_passes(self):
        """Excel's uncomputable marker is text, and text is not 'close enough'."""
        lines = self._compare({14: 100.0}, {14: "-"})
        self.assertEqual(lines[14]["verdict"], "uncomputable")
        self.assertNotEqual(lines[14]["verdict"], "reproduced")

    def test_a_missing_result_fails_rather_than_passes(self):
        lines = self._compare({14: 100.0}, {})
        self.assertEqual(lines[14]["verdict"], "uncomputable")

    def test_a_declared_exemption_excuses_a_deviation(self):
        exemptions = {16: {"reproduction_exemption": "identity_omits_line",
                           "rationale": "the identity omits other income"}}
        lines = self._compare({16: 107787.0}, {16: 82035.71}, exemptions)
        self.assertEqual(lines[16]["verdict"], "exempt")
        self.assertEqual(lines[16]["reproduction_exemption"], "identity_omits_line")
        self.assertIn("other income", lines[16]["rationale"])

    def test_an_exemption_does_not_relabel_a_line_that_already_reproduces(self):
        """An exemption is permission to differ, not a verdict of its own."""
        exemptions = {14: {"reproduction_exemption": "forward_by_design",
                           "rationale": "buybacks"}}
        lines = self._compare({14: 100.0}, {14: 100.0}, exemptions)
        self.assertEqual(lines[14]["verdict"], "reproduced")

    def test_an_exemption_cannot_hide_an_uncomputable_line(self):
        exemptions = {14: {"reproduction_exemption": "forward_by_design",
                           "rationale": "buybacks"}}
        lines = self._compare({14: 100.0}, {14: "-"}, exemptions)
        self.assertEqual(lines[14]["verdict"], "uncomputable")


class ExemptionDeclarationTests(unittest.TestCase):
    """An exemption has to be declared properly to count as declared."""

    PERMITTED = {"forward_by_design", "identity_omits_line", "driver_spans_multiple_years"}

    def test_a_well_formed_exemption_is_accepted(self):
        record = {"base_year_reproduction": {"exemptions": [
            {"row": 19, "reproduction_exemption": "forward_by_design",
             "rationale": "shares are net of buybacks"}]}}
        exemptions, problems = gate.declared_exemptions(record, self.PERMITTED)
        self.assertEqual(problems, [])
        self.assertIn(19, exemptions)

    def test_an_unknown_code_is_rejected(self):
        record = {"base_year_reproduction": {"exemptions": [
            {"row": 19, "reproduction_exemption": "because_i_said_so",
             "rationale": "trust me"}]}}
        exemptions, problems = gate.declared_exemptions(record, self.PERMITTED)
        self.assertEqual(exemptions, {})
        self.assertTrue(any("not a permitted value" in p for p in problems))

    def test_an_exemption_without_a_rationale_is_rejected(self):
        """A code alone records that someone excused a line, not why."""
        record = {"base_year_reproduction": {"exemptions": [
            {"row": 19, "reproduction_exemption": "forward_by_design", "rationale": "  "}]}}
        exemptions, problems = gate.declared_exemptions(record, self.PERMITTED)
        self.assertEqual(exemptions, {})
        self.assertTrue(any("no rationale" in p for p in problems))

    def test_a_case_with_no_block_declares_no_exemptions(self):
        exemptions, problems = gate.declared_exemptions({}, self.PERMITTED)
        self.assertEqual((exemptions, problems), ({}, []))

    def test_the_permitted_codes_come_from_the_glossary_not_from_this_tool(self):
        """Never a second copy of a controlled value set.

        If the tool carried its own list, the glossary and the gate could
        disagree and nothing would notice. This asserts they are the same
        object of knowledge by reading the glossary here independently.
        """
        from_tool = gate.permitted_exemption_codes()
        glossary = json.loads(
            (ROOT / "standards" / "vocabulary" / "glossary.json").read_text(encoding="utf-8")
        )
        from_glossary = set(
            glossary["controlled_fields"]["reproduction_exemption"]["permitted_values"]
        )
        self.assertEqual(from_tool, from_glossary)
        self.assertTrue(from_tool, "the glossary declares no exemption codes at all")


@unittest.skipUnless(shutil.which("soffice"), "LibreOffice is required to evaluate formulas")
class RecalculationTests(unittest.TestCase):
    """End to end, through the same LibreOffice recalculation the repository uses."""

    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp(prefix="byr_test_"))
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.permitted = gate.permitted_exemption_codes()

    def check(self, path: Path, record=None):
        return gate.check_workbook(path, record or {}, self.permitted, 0.5)

    def test_a_sound_chain_reproduces_its_base_year(self):
        path = build_workbook(self.tmp / "sound.xlsx", actuals=SOUND_ACTUALS)
        result = self.check(path)
        self.assertEqual(result["outcome"], "reproduced", result)
        verdicts = {line["row"]: line["verdict"] for line in result["lines"]}
        self.assertEqual(set(verdicts.values()), {"reproduced"})

    def test_a_planted_driver_error_is_caught(self):
        """The control for the test above: the gate must be able to say no.

        A pass is only evidence while a fail is reachable, so the same fixture
        is rebuilt with one driver moved and the gate has to notice.
        """
        path = build_workbook(self.tmp / "broken.xlsx", actuals=SOUND_ACTUALS,
                              opex_driver=0.30)
        result = self.check(path)
        self.assertEqual(result["outcome"], "deviates", result)
        by_row = {line["row"]: line for line in result["lines"]}
        self.assertEqual(by_row[10]["verdict"], "deviates")   # opex
        self.assertEqual(by_row[11]["verdict"], "deviates")   # EBIT below it
        self.assertEqual(by_row[7]["verdict"], "reproduced")  # COGS is untouched

    def test_a_workbook_whose_only_actual_is_the_anchor_is_not_a_pass(self):
        """The vacuous pass this gate shipped with in its first version.

        The probe HOLDS revenue equal to the disclosed figure, so a case that
        sources revenue and nothing else compares the probe with itself and
        can only ever agree.
        """
        path = build_workbook(self.tmp / "anchor.xlsx", actuals={5: 1000.0})
        result = self.check(path)
        self.assertEqual(result["outcome"], "anchor_only", result)
        self.assertNotEqual(result["outcome"], "reproduced")
        self.assertIn("by construction", result["detail"])

    def test_the_anchor_row_is_never_among_the_compared_lines(self):
        path = build_workbook(self.tmp / "full.xlsx", actuals=SOUND_ACTUALS)
        result = self.check(path)
        self.assertNotIn(5, [line["row"] for line in result["lines"]])

    def test_an_empty_actuals_column_is_reported_not_passed(self):
        path = build_workbook(self.tmp / "empty.xlsx", actuals={})
        result = self.check(path)
        self.assertEqual(result["outcome"], "no_actuals")
        self.assertNotEqual(result["outcome"], "reproduced")

    def test_a_workbook_without_a_statement_sheet_is_not_applicable(self):
        path = build_workbook(self.tmp / "nois.xlsx", include_is=False)
        self.assertEqual(self.check(path)["outcome"], "not_applicable")

    def test_a_malformed_exemption_fails_the_workbook_rather_than_being_ignored(self):
        path = build_workbook(self.tmp / "bad_exempt.xlsx", actuals=SOUND_ACTUALS)
        record = {"base_year_reproduction": {"exemptions": [
            {"row": 11, "reproduction_exemption": "invented_code", "rationale": "x"}]}}
        result = self.check(path, record)
        self.assertEqual(result["outcome"], "unreadable")
        self.assertIn("not a permitted value", result["detail"])
