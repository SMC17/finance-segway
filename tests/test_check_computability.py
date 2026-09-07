"""Guardrails for the check-computability gate.

The defect this gate exists to stop is inverted, not merely silent: an
uncomputable metric does not fail a threshold test, it wins it. Excel sorts
text above every number, so `IF("-">=3,"PASS",...)` awards the top band,
and `IFERROR(a/b, 0)` hands 0 to a test where low is good. Both were live
in the committed corpus.

So the tests below are weighted toward discrimination. A gate that flagged
every non-numeric metric would pass the negative cases here and be useless,
because several checks are deliberately text-valued -- "SUPPORTED",
"NONE" -- and one of those correctly reports FAIL.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tools import verify_check_computability as gate


def _workbook(rows, *, sheet="Decision & Checks", status_col=4, metric_col=3) -> Path:
    """A minimal check sheet. Literal values, no formulas, so the cached
    and formula views agree and no recalculation is needed."""
    book = Workbook()
    ws = book.active
    ws.title = sheet
    ws.cell(row=4, column=2, value="Check / decision")
    for i, (label, metric, status) in enumerate(rows, start=5):
        ws.cell(row=i, column=2, value=label)
        if metric is not None:
            ws.cell(row=i, column=metric_col, value=metric)
        ws.cell(row=i, column=status_col, value=status)
    path = Path(tempfile.mkdtemp()) / "case.xlsx"
    book.save(path)
    return path


class CheckComputabilityTests(unittest.TestCase):
    def test_a_pass_on_the_uncomputable_marker_is_caught(self) -> None:
        path = _workbook([
            ("Leverage ratio", 0.4, "PASS"),
            ("LTV / CAC", "-", "PASS"),
            ("Coverage", 3.1, "PASS"),
        ])
        rows = gate.inspect(path)["rows"]
        bad = [r for r in rows if r["state"] == "passes_on_uncomputable"]
        self.assertEqual(1, len(bad))
        self.assertEqual("LTV / CAC", bad[0]["label"])

    def test_the_marker_reported_honestly_is_not_a_violation(self) -> None:
        """The repository's own convention: uncomputable becomes REVIEW.
        Sixteen checks already do this and must not be flagged."""
        path = _workbook([
            ("Roll yield", "-", "REVIEW"),
            ("Fulcrum security", "-", "REVIEW"),
            ("Coverage", 3.1, "PASS"),
        ])
        rows = gate.inspect(path)["rows"]
        self.assertEqual(0, sum(1 for r in rows if r["state"] == "passes_on_uncomputable"))
        self.assertEqual(2, sum(1 for r in rows if r["state"] == "uncomputable_reported"))

    def test_a_deliberately_text_valued_check_is_left_alone(self) -> None:
        """Discrimination. Flagging every non-numeric metric would catch
        these too, and they are correct: a check whose metric IS a word."""
        path = _workbook([
            ("Liquidation terms supported", "SUPPORTED", "PASS"),
            ("New-money commitment status", "PASS", "PASS"),
            ("Holder-election equilibrium", "NONE", "FAIL"),
        ])
        rows = gate.inspect(path)["rows"]
        self.assertEqual(0, sum(1 for r in rows if r["state"] != "ok"))

    def test_a_zero_in_a_currency_format_is_not_the_marker(self) -> None:
        """The control that makes the marker meaningful.

        The repo's currency format displays zero as "-". If openpyxl read
        the DISPLAY rather than the VALUE, every zero metric would look
        uncomputable and this gate would be noise. It reads the value.
        """
        path = _workbook([("Residual", 0, "PASS"), ("Coverage", 3.1, "PASS"), ("Leverage", 0.4, "PASS")])
        book = load_workbook(path)
        book["Decision & Checks"].cell(row=5, column=3).number_format = '$#,##0;($#,##0);"-"'
        book.save(path)
        self.assertEqual(0, load_workbook(path, data_only=True)["Decision & Checks"]["C5"].value)
        rows = gate.inspect(path)["rows"]
        self.assertEqual(["ok", "ok", "ok"], [r["state"] for r in rows])

    def test_the_status_column_is_located_not_assumed(self) -> None:
        """Two archetypes put the verdict in different columns. A gate that
        assumed one would silently inspect nothing in the other -- which is
        exactly how the first version of this analysis reported 30 of 54
        workbooks as having no checks at all."""
        four_col = _workbook([("A", 1, "PASS"), ("B", 2, "PASS"), ("C", 3, "PASS")], status_col=4)
        two_col = _workbook([("A", None, "PASS"), ("B", None, "PASS"), ("C", None, "PASS")],
                            sheet="Checks", status_col=3)
        self.assertEqual(3, len(gate.inspect(four_col)["rows"]))
        rows = gate.inspect(two_col)["rows"]
        self.assertEqual(3, len(rows))
        self.assertTrue(all(r["state"] == "no_metric_shown" for r in rows))

    def test_a_sheet_with_no_metric_column_is_declared_not_passed(self) -> None:
        """"Cannot be asked" must not render as "asked and fine"."""
        path = _workbook([("A", None, "PASS"), ("B", None, "PASS"), ("C", None, "PASS")],
                         sheet="Checks", status_col=3)
        rows = gate.inspect(path)["rows"]
        self.assertEqual(0, sum(1 for r in rows if r["state"] == "ok"))
        self.assertEqual(3, sum(1 for r in rows if r["state"] == "no_metric_shown"))

    def test_a_workbook_it_cannot_read_is_counted_not_skipped(self) -> None:
        """The status column is found by looking for three or more verdicts.
        A workbook with fewer is not inspectable -- and must be reported as
        such rather than contributing a silent zero to a green total."""
        path = _workbook([("Only one check", 1, "PASS")])
        result = gate.inspect(path)
        self.assertEqual("no_status_column", result["state"])
        self.assertEqual([], result["rows"])

    def test_the_committed_corpus_passes(self) -> None:
        report = gate.assess()
        self.assertEqual("PASS", report["status"], report["violations"])
        self.assertEqual(0, report["counts"]["passes_on_uncomputable"])

    def test_the_corpus_scan_is_not_vacuous(self) -> None:
        """A gate over zero checks passes. Pin that it really inspected a
        population, and that the honest-marker convention is present -- so
        a change that silently emptied the scan reddens here."""
        report = gate.assess()
        self.assertGreaterEqual(report["counts"]["workbooks"], 50)
        self.assertGreaterEqual(report["counts"]["checks_inspected"], 150)
        self.assertGreaterEqual(report["counts"]["uncomputable_reported_honestly"], 10)


if __name__ == "__main__":
    unittest.main()
