"""Require a model to reproduce the period it was calibrated on.

WHY THIS GATE EXISTS
--------------------
Every other check in this repository asks whether a workbook is internally
consistent: that it recalculates without errors, that its identities close,
that its inputs are sourced, that its vocabulary is controlled. A model can
pass all of them and still be wrong about the world, because none of them ever
compares the model to an outcome.

This one does the cheapest external comparison available. Where a case carries
disclosed actuals for its base year, hold revenue at the disclosed base-year
figure, leave every driver exactly as committed, recalculate, and compare each
forecast line to the actual beside it. A chain fed its own base year must
return its own base year. When it does not, the model disagrees with a fact its
own author recorded one column to the left.

It found the defect it was written for: the Microsoft corporate-finance case
reproduces revenue, COGS, opex and interest expense to the cent and misses EBIT
by 24,462 of 109,433 -- 22.35% -- because the opex driver is derived from
`OperatingIncomeLoss`, which is already net of depreciation and amortisation,
so the row labelled EBITDA computes EBIT and the next row subtracts D&A again.

NOT EVERY LINE SHOULD REPRODUCE, AND THAT IS THE HARD PART
----------------------------------------------------------
Two lines in that same case legitimately differ from the base year:

  * Diluted shares are a forward count net of buybacks, so FY1 *should* sit
    below FY0. Reproducing the base year here would be the error.
  * Pre-tax income, tax and net income sit 1.20% low because the template's
    forward identity is `pre-tax = EBIT - interest` and has no row for other
    income and expense, which the disclosure includes.

Both are recorded in the case's prose today. Prose is not something a gate can
read, so this tool requires the exemption to be declared in the case record as
a `base_year_reproduction.exemptions` entry carrying a controlled
`reproduction_exemption` code and a rationale. An undeclared deviation fails;
a declared one is printed with its reason and does not. The effect is that
choosing not to reproduce a line becomes a reviewable statement in the record
rather than a silence in the arithmetic.

WHAT THE OUTCOMES MEAN
----------------------
  reproduced      every non-exempt line is within tolerance
  deviates        at least one non-exempt line is outside it            -> FAIL
  no_actuals      the base-year column is empty: nothing to compare to
  no_base         the base-year revenue is absent or zero
  anchor_only     revenue is the only disclosed actual, and this probe holds
                  revenue equal by construction, so nothing here can disagree
  unreadable      the workbook or its record could not be read          -> FAIL
  recalc_failed   LibreOffice could not evaluate the probe              -> FAIL

`no_actuals`, `no_base` and `anchor_only` are not passes. A case that cannot be checked and a
case that reproduces perfectly must never render as the same green, so they are
counted and named separately, and a run in which *nothing* could be checked
exits 3 rather than 0 -- because "no model in this repository can be compared
to a period it can be verified on" is a finding, not a clean bill.

Usage:
    python tools/verify_base_year_reproduction.py
    python tools/verify_base_year_reproduction.py --report base-year.json
    python tools/verify_base_year_reproduction.py --tolerance-pct 0.5
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

from recalc import recalc  # noqa: E402

CASE_DIR = ROOT / "standards" / "public_cases"
GLOSSARY = ROOT / "standards" / "vocabulary" / "glossary.json"

STATEMENT_SHEET = "IS"
ACTUAL_HEADER = re.compile(r"^FY.*A$")
ESTIMATE_HEADER = re.compile(r"^FY\d+E$")
HEADER_SEARCH_ROWS = 8
LINE_SEARCH_ROWS = 30
LABEL_COLUMN = 2
DEFAULT_TOLERANCE_PCT = 0.5

FAIL_OUTCOMES = {"deviates", "unreadable", "recalc_failed"}
EXIT_OK = 0
EXIT_FAILED = 1
EXIT_NOTHING_CHECKED = 3


def permitted_exemption_codes() -> set[str]:
    """Read the codes from the glossary rather than restating them here.

    A second copy of a controlled value set is how the two copies diverge, so
    the vocabulary lives in one file and this tool asks it.

    Raises rather than returning an empty set when the glossary cannot answer.
    An empty permitted set would make every exemption code acceptable, so a
    missing or emptied vocabulary would silently turn the strictest part of
    this gate into a rubber stamp -- the failure would look like a pass.
    """
    if not GLOSSARY.exists():
        raise FileNotFoundError(
            f"missing {GLOSSARY}: the permitted exemption codes cannot be read, and "
            "an unreadable vocabulary must not be treated as permission"
        )
    glossary = json.loads(GLOSSARY.read_text(encoding="utf-8"))
    field = glossary.get("controlled_fields", {}).get("reproduction_exemption", {})
    codes = set(field.get("permitted_values", {}))
    if not codes:
        raise ValueError(
            "the glossary declares no permitted values for reproduction_exemption; "
            "refusing to accept every code by default"
        )
    return codes


def find_columns(sheet: Any) -> tuple[int, int, int] | None:
    """Locate the header row, the base-year actual column and the first estimate."""
    for row in range(1, HEADER_SEARCH_ROWS + 1):
        actuals, estimates = [], []
        for column in range(3, 14):
            value = sheet.cell(row=row, column=column).value
            if not isinstance(value, str):
                continue
            if ACTUAL_HEADER.match(value):
                actuals.append(column)
            elif ESTIMATE_HEADER.match(value):
                estimates.append(column)
        if actuals and estimates:
            return row, max(actuals), min(estimates)
    return None


def disclosed_actuals(sheet: Any, header_row: int,
                      base_column: int) -> tuple[dict[int, float], list[int]]:
    """The rows whose base-year cell carries a literal number, and those skipped.

    Only literals count. A base-year cell holding a formula is a line the model
    DERIVES rather than a figure anyone sourced, and comparing a derived
    historical line to a derived forecast line largely restates the arithmetic
    that produced both. The skipped rows are returned so the narrowing can be
    printed rather than left for a reader to infer from a shorter table.
    """
    found: dict[int, float] = {}
    skipped: list[int] = []
    for row in range(header_row + 1, header_row + LINE_SEARCH_ROWS):
        value = sheet.cell(row=row, column=base_column).value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            found[row] = float(value)
        elif value is not None:
            skipped.append(row)
    return found, skipped


def revenue_row(sheet: Any, header_row: int) -> int | None:
    for row in range(header_row + 1, header_row + LINE_SEARCH_ROWS):
        label = sheet.cell(row=row, column=LABEL_COLUMN).value
        if isinstance(label, str) and label.strip().lower() == "revenue":
            return row
    return None


def run_at_base_year(workbook_path: Path, header_row: int, base_column: int,
                     estimate_column: int, anchor_row: int) -> dict[int, Any]:
    """Hold revenue at the disclosed base year, recalculate, read the column back.

    Only the revenue anchor is touched. Every driver keeps the value it was
    committed with, so what comes back is the committed chain evaluated at a
    period whose answer is already known.
    """
    with tempfile.TemporaryDirectory(prefix="base_year_") as workspace:
        scratch = Path(workspace) / workbook_path.name
        shutil.copy(workbook_path, scratch)
        book = openpyxl.load_workbook(scratch)
        sheet = book[STATEMENT_SHEET]
        sheet.cell(row=anchor_row, column=estimate_column).value = (
            f"={get_column_letter(base_column)}{anchor_row}"
        )
        book.save(scratch)
        recalc(str(scratch), force=True)
        evaluated = openpyxl.load_workbook(scratch, data_only=True)[STATEMENT_SHEET]
        return {
            row: evaluated.cell(row=row, column=estimate_column).value
            for row in range(header_row + 1, header_row + LINE_SEARCH_ROWS)
        }


def case_records() -> dict[str, dict[str, Any]]:
    records: dict[str, dict[str, Any]] = {}
    if not CASE_DIR.is_dir():
        return records
    for path in sorted(CASE_DIR.glob("*.json")):
        try:
            record = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            continue
        output = record.get("output")
        if isinstance(output, str):
            records[output] = record
    return records


def declared_exemptions(record: dict[str, Any], permitted: set[str]) -> tuple[dict[int, dict], list[str]]:
    """Exemptions the case declares, and complaints about any that are malformed."""
    block = record.get("base_year_reproduction") or {}
    exemptions, problems = {}, []
    for entry in block.get("exemptions", []):
        row, code = entry.get("row"), entry.get("reproduction_exemption")
        rationale = (entry.get("rationale") or "").strip()
        if not isinstance(row, int):
            problems.append(f"exemption without an integer row: {entry!r}")
            continue
        if code not in permitted:
            problems.append(f"row {row}: reproduction_exemption {code!r} is not a permitted value")
            continue
        if not rationale:
            problems.append(f"row {row}: exemption carries no rationale")
            continue
        exemptions[row] = entry
    return exemptions, problems


def compare(disclosed: dict[int, float], reproduced: dict[int, Any], labels: dict[int, str],
            exemptions: dict[int, dict], tolerance_pct: float) -> list[dict[str, Any]]:
    lines = []
    for row in sorted(disclosed):
        want = disclosed[row]
        got = reproduced.get(row)
        entry: dict[str, Any] = {
            "row": row,
            "line": labels.get(row, ""),
            "disclosed": want,
            "reproduced": got,
        }
        if not isinstance(got, (int, float)) or isinstance(got, bool):
            entry["verdict"] = "uncomputable"
        elif want == 0:
            entry["verdict"] = "reproduced" if got == 0 else "deviates"
            entry["delta"] = float(got) - want
        else:
            delta = float(got) - want
            entry["delta"] = delta
            entry["delta_pct"] = delta / abs(want) * 100.0
            entry["verdict"] = (
                "reproduced" if abs(entry["delta_pct"]) <= tolerance_pct else "deviates"
            )
        # An exemption says a line is expected to DIFFER from the base year. It
        # says nothing about a line that failed to compute at all, and must not
        # be allowed to excuse one -- an uncomputable line is a broken model,
        # not a modelled difference.
        if row in exemptions and entry["verdict"] == "deviates":
            entry["verdict"] = "exempt"
            entry["reproduction_exemption"] = exemptions[row].get("reproduction_exemption")
            entry["rationale"] = exemptions[row].get("rationale")
        lines.append(entry)
    return lines


def repo_relative(path: Path) -> str:
    """Name a path relative to the repository when it is inside it.

    A workbook under test may live in a temporary directory, and a report is
    not worth raising a ValueError over.
    """
    try:
        return path.resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return path.as_posix()


def check_workbook(path: Path, record: dict[str, Any], permitted: set[str],
                   tolerance_pct: float) -> dict[str, Any]:
    relative = repo_relative(path)
    result: dict[str, Any] = {"workbook": relative, "lines": []}
    try:
        book = openpyxl.load_workbook(path)
    except Exception as error:  # openpyxl raises a wide family on damaged files
        return {**result, "outcome": "unreadable", "detail": str(error)[:200]}
    if STATEMENT_SHEET not in book.sheetnames:
        return {**result, "outcome": "not_applicable",
                "detail": f"no {STATEMENT_SHEET} sheet: not a BASE-template statement model"}
    sheet = book[STATEMENT_SHEET]
    columns = find_columns(sheet)
    if columns is None:
        return {**result, "outcome": "not_applicable",
                "detail": "no actual/estimate column headers"}
    header_row, base_column, estimate_column = columns
    if base_column >= estimate_column:
        return {**result, "outcome": "unreadable",
                "detail": f"the base-year column {get_column_letter(base_column)} is not "
                          f"left of the first forecast column "
                          f"{get_column_letter(estimate_column)}; the layout is not the "
                          "one this probe understands"}
    result["base_column"] = get_column_letter(base_column)
    result["estimate_column"] = get_column_letter(estimate_column)
    anchor = revenue_row(sheet, header_row)
    if anchor is None:
        return {**result, "outcome": "not_applicable", "detail": "no revenue line to anchor"}
    disclosed, derived_rows = disclosed_actuals(sheet, header_row, base_column)
    result["derived_base_year_rows"] = derived_rows
    if not disclosed:
        return {**result, "outcome": "no_actuals",
                "detail": "the base-year column is empty, so the model cannot be "
                          "compared to any period it can be verified on"}
    if disclosed.get(anchor) in (None, 0):
        return {**result, "outcome": "no_base",
                "detail": "the base-year revenue is absent or zero"}
    # The anchor is HELD equal to the disclosed figure by this tool, so comparing
    # it compares the probe to itself. Dropping it is what stops a case whose only
    # sourced actual is revenue from reporting a pass it did not earn.
    comparable = {row: value for row, value in disclosed.items() if row != anchor}
    if not comparable:
        return {**result, "outcome": "anchor_only",
                "detail": f"the only disclosed actual is the revenue anchor on row "
                          f"{anchor}, which this probe holds equal by construction; "
                          "there is nothing here that could disagree"}
    exemptions, problems = declared_exemptions(record, permitted)
    if problems:
        return {**result, "outcome": "unreadable",
                "detail": "; ".join(problems)}
    try:
        reproduced = run_at_base_year(path, header_row, base_column, estimate_column, anchor)
    except Exception as error:
        return {**result, "outcome": "recalc_failed", "detail": str(error)[:200]}
    labels = {
        row: sheet.cell(row=row, column=LABEL_COLUMN).value or ""
        for row in comparable
    }
    result["lines"] = compare(comparable, reproduced, labels, exemptions, tolerance_pct)
    verdicts = {line["verdict"] for line in result["lines"]}
    result["outcome"] = "deviates" if ("deviates" in verdicts or "uncomputable" in verdicts) else "reproduced"
    return result


def candidate_workbooks() -> list[Path]:
    return sorted(
        path for path in ROOT.glob("*/instances/*.xlsx")
        if not path.name.startswith("~$")
    )


def render(results: list[dict[str, Any]], tolerance_pct: float) -> None:
    for result in results:
        outcome = result["outcome"]
        if outcome == "not_applicable":
            continue
        print(f"\n{result['workbook']}  [{outcome}]")
        if result.get("detail"):
            print(f"    {result['detail']}")
        for line in result["lines"]:
            mark = {"reproduced": "  ok ", "deviates": " FAIL", "exempt": " exmp",
                    "uncomputable": " FAIL"}[line["verdict"]]
            disclosed = f"{line['disclosed']:,.2f}"
            got = line["reproduced"]
            got_text = f"{got:,.2f}" if isinstance(got, (int, float)) else repr(got)
            tail = ""
            if "delta_pct" in line:
                tail = f"  delta={line['delta']:>+13,.2f} ({line['delta_pct']:+.2f}%)"
            print(f"{mark} r{line['row']:<3} {str(line['line'])[:24]:24s}"
                  f" disclosed={disclosed:>15s}  reproduced={got_text:>15s}{tail}")
            if line["verdict"] == "exempt":
                print(f"        exempt [{line.get('reproduction_exemption')}]: "
                      f"{line.get('rationale')}")

    counts: dict[str, int] = {}
    for result in results:
        counts[result["outcome"]] = counts.get(result["outcome"], 0) + 1
    print("\n" + "-" * 74)
    print(f"tolerance: {tolerance_pct}% of the disclosed figure")
    for outcome in sorted(counts):
        print(f"  {outcome:16s} {counts[outcome]}")

    derived = sum(len(r.get("derived_base_year_rows") or []) for r in results)
    print("\nAND HERE IS WHAT THIS DID NOT CHECK")
    print("  - Only the income statement, and only lines whose base-year cell carries")
    print("    a literal number. A line the case never sourced is invisible here.")
    if derived:
        print(f"  - {derived} base-year cell(s) hold a FORMULA rather than a sourced")
        print("    figure and are skipped: comparing a derived historical line to a")
        print("    derived forecast line largely restates the arithmetic behind both.")
    print("  - Only the FIRST forecast column. A driver that fades across later columns")
    print("    is exercised at its first value alone.")
    print("  - Reproducing the base year does not make a forecast right. It establishes")
    print("    that the chain does not contradict the one period whose answer is known.")
    print("  - The disclosed figures themselves are taken as given. This compares the")
    print("    model to the case's own recorded actuals, not to the filing.")
    not_checked = (counts.get("no_actuals", 0) + counts.get("no_base", 0)
                   + counts.get("anchor_only", 0))
    if not_checked:
        print(f"  - {not_checked} workbook(s) carry the structure but no comparable period,")
        print("    and are reported above rather than counted as passes.")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    parser.add_argument("--report", type=Path, help="write the full result as JSON")
    parser.add_argument("--tolerance-pct", type=float, default=DEFAULT_TOLERANCE_PCT,
                        help=f"allowed deviation, percent (default {DEFAULT_TOLERANCE_PCT})")
    args = parser.parse_args()

    permitted = permitted_exemption_codes()
    records = case_records()
    results = []
    for path in candidate_workbooks():
        relative = repo_relative(path)
        results.append(check_workbook(path, records.get(relative, {}), permitted,
                                      args.tolerance_pct))
    render(results, args.tolerance_pct)

    if args.report:
        args.report.write_text(
            json.dumps({"tolerance_pct": args.tolerance_pct, "results": results},
                       indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        print(f"\nreport: {args.report}")

    failed = [r for r in results if r["outcome"] in FAIL_OUTCOMES]
    checked = [r for r in results if r["outcome"] in {"reproduced", "deviates"}]
    if failed:
        print(f"\nFAIL: {len(failed)} workbook(s) do not reproduce the period they were "
              "calibrated on.")
        return EXIT_FAILED
    if not checked:
        print("\nFAIL: no workbook in this repository could be compared to a period it "
              "can be verified on. That is a finding, not a clean bill.")
        return EXIT_NOTHING_CHECKED
    print(f"\nPASS: {len(checked)} workbook(s) reproduce their base year.")
    return EXIT_OK


if __name__ == "__main__":
    raise SystemExit(main())
